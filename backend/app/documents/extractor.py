import os
from typing import List, Dict, Any
from pathlib import Path
from pypdf import PdfReader
from docx import Document as DocxDocument
from app.documents.cleaner import clean_text

class ExtractedPage:
    def __init__(self, page_number: int, text: str):
        self.page_number = page_number
        self.text = text

class DocumentExtractor:
    @staticmethod
    def extract_file(file_path: str) -> List[ExtractedPage]:
        """Extract pages/sections from PDF, DOCX, TXT, or MD files."""
        path = Path(file_path)
        ext = path.suffix.lower()

        if not path.exists():
            raise FileNotFoundError(f"File not found: {file_path}")

        if ext == ".pdf":
            return DocumentExtractor._extract_pdf(file_path)
        elif ext in [".docx", ".doc"]:
            return DocumentExtractor._extract_docx(file_path)
        elif ext in [".txt", ".md", ".markdown"]:
            return DocumentExtractor._extract_txt(file_path)
        else:
            raise ValueError(f"Unsupported file format: {ext}. Supported formats: .pdf, .docx, .txt, .md")

    @staticmethod
    def _extract_pdf(file_path: str) -> List[ExtractedPage]:
        pages = []
        try:
            reader = PdfReader(file_path)
            for idx, page in enumerate(reader.pages, start=1):
                raw_text = page.extract_text() or ""
                cleaned = clean_text(raw_text)
                if cleaned:
                    pages.append(ExtractedPage(page_number=idx, text=cleaned))
        except Exception as e:
            raise RuntimeError(f"Failed to extract text from PDF: {str(e)}")
        return pages

    @staticmethod
    def _extract_docx(file_path: str) -> List[ExtractedPage]:
        try:
            doc = DocxDocument(file_path)
            text_chunks = []
            for para in doc.paragraphs:
                cleaned = clean_text(para.text)
                if cleaned:
                    text_chunks.append(cleaned)
            for table in doc.tables:
                for row in table.rows:
                    row_text = " | ".join(clean_text(cell.text) for cell in row.cells if cell.text.strip())
                    if row_text:
                        text_chunks.append(row_text)
            
            full_text = "\n\n".join(text_chunks)
            return [ExtractedPage(page_number=1, text=full_text)]
        except Exception as e:
            raise RuntimeError(f"Failed to extract text from DOCX: {str(e)}")

    @staticmethod
    def _extract_txt(file_path: str) -> List[ExtractedPage]:
        try:
            with open(file_path, "r", encoding="utf-8", errors="replace") as f:
                content = f.read()
            cleaned = clean_text(content)
            # Break large text files into pages of ~2000 chars if long
            page_size = 2000
            if len(cleaned) <= page_size:
                return [ExtractedPage(page_number=1, text=cleaned)]
            
            paragraphs = cleaned.split("\n\n")
            pages = []
            curr_text = []
            curr_len = 0
            page_num = 1
            for p in paragraphs:
                curr_text.append(p)
                curr_len += len(p)
                if curr_len >= page_size:
                    pages.append(ExtractedPage(page_number=page_num, text="\n\n".join(curr_text)))
                    curr_text = []
                    curr_len = 0
                    page_num += 1
            if curr_text:
                pages.append(ExtractedPage(page_number=page_num, text="\n\n".join(curr_text)))
            return pages
        except Exception as e:
            raise RuntimeError(f"Failed to extract text from file: {str(e)}")
